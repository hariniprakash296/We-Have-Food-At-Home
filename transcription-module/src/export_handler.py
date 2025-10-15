"""
Export handler for generating transcripts in multiple formats.
Supports TXT, DOCX, PDF, and JSON exports.
"""

import json
from pathlib import Path
from datetime import datetime
from typing import Optional
import logging

try:
    from docx import Document
    from docx.shared import Pt, Inches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logging.warning("python-docx not available. DOCX export disabled.")

try:
    from fpdf import FPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logging.warning("fpdf2 not available. PDF export disabled.")

from config import get_settings
from transcription_engine import TranscriptionSession

logger = logging.getLogger(__name__)


class ExportHandler:
    """
    Handles exporting transcripts to various formats.
    
    Supported formats:
    - TXT: Plain text
    - DOCX: Microsoft Word document
    - PDF: Portable Document Format
    - JSON: Structured data format
    """
    
    def __init__(self):
        """Initialize the export handler."""
        self.settings = get_settings()
        self.export_dir = Path(self.settings.export_directory)
        self.export_dir.mkdir(exist_ok=True, parents=True)
    
    def export(
        self,
        session: TranscriptionSession,
        filename: str,
        format: str = "txt"
    ):
        """
        Export transcript to specified format.
        
        Args:
            session: Transcription session to export
            filename: Output filename
            format: Export format (txt, docx, pdf, json)
        """
        # Ensure file is in export directory
        filepath = self.export_dir / filename
        
        # Route to appropriate exporter
        exporters = {
            'txt': self._export_txt,
            'docx': self._export_docx,
            'pdf': self._export_pdf,
            'json': self._export_json
        }
        
        exporter = exporters.get(format.lower())
        if not exporter:
            raise ValueError(f"Unsupported export format: {format}")
        
        exporter(session, filepath)
        logger.info(f"Exported to {filepath}")
    
    def _export_txt(
        self,
        session: TranscriptionSession,
        filepath: Path
    ):
        """Export to plain text format."""
        with open(filepath, 'w', encoding='utf-8') as f:
            # Write header
            f.write(f"Meeting Transcript\n")
            f.write(f"Session ID: {session.session_id}\n")
            f.write(f"Date: {session.start_time.strftime('%Y-%m-%d %H:%M:%S')}\n")
            f.write(f"{'=' * 60}\n\n")
            
            # Write segments
            for segment in session.segments:
                timestamp = datetime.fromtimestamp(segment.timestamp).strftime('%H:%M:%S')
                speaker = f"[{segment.speaker_id}] " if segment.speaker_id else ""
                f.write(f"[{timestamp}] {speaker}{segment.text}\n")
            
            # Write footer
            f.write(f"\n{'=' * 60}\n")
            f.write(f"Total segments: {len(session.segments)}\n")
    
    def _export_docx(
        self,
        session: TranscriptionSession,
        filepath: Path
    ):
        """Export to Microsoft Word format."""
        if not DOCX_AVAILABLE:
            logger.error("DOCX export not available. Install python-docx.")
            return
        
        doc = Document()
        
        # Add title
        title = doc.add_heading('Meeting Transcript', 0)
        
        # Add metadata
        doc.add_paragraph(f"Session ID: {session.session_id}")
        doc.add_paragraph(f"Date: {session.start_time.strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()
        
        # Add transcript
        for segment in session.segments:
            timestamp = datetime.fromtimestamp(segment.timestamp).strftime('%H:%M:%S')
            speaker = f"{segment.speaker_id}: " if segment.speaker_id else ""
            
            # Add paragraph with formatting
            para = doc.add_paragraph()
            
            # Add timestamp in italic
            run_time = para.add_run(f"[{timestamp}] ")
            run_time.italic = True
            run_time.font.size = Pt(10)
            
            # Add speaker in bold (if present)
            if speaker:
                run_speaker = para.add_run(speaker)
                run_speaker.bold = True
            
            # Add text
            para.add_run(segment.text)
        
        # Add footer
        doc.add_paragraph()
        footer = doc.add_paragraph(f"Total segments: {len(session.segments)}")
        footer.runs[0].italic = True
        
        # Save document
        doc.save(str(filepath))
    
    def _export_pdf(
        self,
        session: TranscriptionSession,
        filepath: Path
    ):
        """Export to PDF format."""
        if not PDF_AVAILABLE:
            logger.error("PDF export not available. Install fpdf2.")
            return
        
        pdf = FPDF()
        pdf.add_page()
        
        # Title
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(0, 10, 'Meeting Transcript', ln=True, align='C')
        pdf.ln(5)
        
        # Metadata
        pdf.set_font('Arial', '', 10)
        pdf.cell(0, 5, f"Session ID: {session.session_id}", ln=True)
        pdf.cell(0, 5, f"Date: {session.start_time.strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        pdf.ln(5)
        
        # Transcript
        pdf.set_font('Arial', '', 11)
        for segment in session.segments:
            timestamp = datetime.fromtimestamp(segment.timestamp).strftime('%H:%M:%S')
            speaker = f"{segment.speaker_id}: " if segment.speaker_id else ""
            
            # Format line
            line = f"[{timestamp}] {speaker}{segment.text}"
            
            # Handle long lines
            pdf.multi_cell(0, 5, line)
            pdf.ln(2)
        
        # Footer
        pdf.ln(5)
        pdf.set_font('Arial', 'I', 9)
        pdf.cell(0, 5, f"Total segments: {len(session.segments)}", ln=True)
        
        # Save PDF
        pdf.output(str(filepath))
    
    def _export_json(
        self,
        session: TranscriptionSession,
        filepath: Path
    ):
        """Export to JSON format."""
        data = {
            'session_id': session.session_id,
            'start_time': session.start_time.isoformat(),
            'metadata': session.metadata,
            'segments': [
                {
                    'text': seg.text,
                    'timestamp': seg.timestamp,
                    'speaker_id': seg.speaker_id,
                    'confidence': seg.confidence,
                    'start_time': seg.start_time,
                    'end_time': seg.end_time
                }
                for seg in session.segments
            ],
            'statistics': {
                'total_segments': len(session.segments),
                'total_duration': (
                    session.segments[-1].timestamp - session.segments[0].timestamp
                    if session.segments else 0
                )
            }
        }
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
    
    def list_exports(self) -> list[str]:
        """
        List all exported files.
        
        Returns:
            List of export filenames
        """
        if not self.export_dir.exists():
            return []
        
        files = []
        for ext in ['*.txt', '*.docx', '*.pdf', '*.json']:
            files.extend([f.name for f in self.export_dir.glob(ext)])
        
        return sorted(files)
